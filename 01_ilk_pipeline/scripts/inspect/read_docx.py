from docx import Document
import os

KLASOR = r"D:\_Development\Projects\UYSM_Projects\data\source\UYSM-EXECELL HESAPLAMALARI 0-300 DOSYA"
for f in os.listdir(KLASOR):
    if f.endswith(".docx"):
        print("Dosya:", f)
        doc = Document(os.path.join(KLASOR, f))
        for para in doc.paragraphs:
            if para.text.strip():
                print(para.text)
        for tablo in doc.tables:
            print("\n[TABLO]")
            for row in tablo.rows:
                satirlar = [cell.text.strip() for cell in row.cells]
                if any(satirlar):
                    print(" | ".join(satirlar))
