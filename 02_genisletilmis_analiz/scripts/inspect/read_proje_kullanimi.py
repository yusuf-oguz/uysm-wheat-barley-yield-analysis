import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

path = r"D:\_Development\Projects\UYSM_Project_2\uysm_flash_bellek_degistirilmemis\hansay 2016 verim\Verim 8\PROJE KULLANIMI HAKKINDA BİLGİLER.docx"
doc = Document(path)

print(f"Toplam paragraf: {len(doc.paragraphs)}")
print(f"Toplam tablo: {len(doc.tables)}")
print("=" * 80)

for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    style = para.style.name
    if text:
        print(f"[{i:04d}] [{style}] {text}")

print("\n" + "=" * 80)
print("TABLOLAR:")
for t_idx, table in enumerate(doc.tables):
    print(f"\n--- Tablo {t_idx+1} ({len(table.rows)} satir x {len(table.columns)} sutun) ---")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.strip().replace('\n', ' | ') for c in row.cells]
        if any(cells):
            print(f"  Satir {r_idx}: {cells}")
